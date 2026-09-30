"""Cost Comparison multi-scenario export (Excel workbook and Tableau .hyper)."""

from __future__ import annotations

import io
from pathlib import Path

import duckdb
import pytest
from openpyxl import load_workbook

from planalign_api import config as api_config
from planalign_api.main import create_app
from planalign_api.models.scenario import ScenarioCreate
from planalign_api.models.workspace import WorkspaceCreate
from planalign_api.routers import analytics as analytics_router
from planalign_api.services import comparison_export_service as export_module
from planalign_api.services.comparison_export_service import (
    ComparisonExportService,
    ExportScenario,
    unique_sheet_names,
)
from planalign_api.services.database_path_resolver import ResolvedDatabasePath
from planalign_api.storage.workspace_storage import WorkspaceStorage

pytestmark = [pytest.mark.fast]

ENDPOINT = "/api/workspaces/{workspace_id}/analytics/compare/export"


def _build_db(
    path: Path, employees: int, years: tuple[int, ...] = (2025, 2026)
) -> Path:
    conn = duckdb.connect(str(path))
    try:
        conn.execute(
            """
            CREATE TABLE fct_workforce_snapshot (
                scenario_id VARCHAR DEFAULT 'default',
                employee_id VARCHAR,
                simulation_year INTEGER,
                current_compensation DECIMAL(12, 2),
                snapshot_created_at TIMESTAMPTZ
            );
            CREATE TABLE stg_census_data (
                employee_id VARCHAR,
                employee_hire_date DATE
            );
            """
        )
        for year in years:
            conn.execute(
                "INSERT INTO fct_workforce_snapshot SELECT 'default', 'e' || i, ?, 50000 + i, "
                "TIMESTAMPTZ '2026-01-01 12:00:00+00' FROM range(?) t(i)",
                [year, employees],
            )
        conn.execute(
            "INSERT INTO stg_census_data SELECT 'e' || i, DATE '2020-01-01' "
            "FROM range(?) t(i)",
            [employees],
        )
    finally:
        conn.close()
    return path


def _scenario(
    tmp_path: Path, name: str, employees: int, census: str = "census.parquet"
) -> ExportScenario:
    return ExportScenario(
        scenario_id=f"id-{name}",
        name=name,
        database_path=_build_db(tmp_path / f"{abs(hash(name))}.duckdb", employees),
        census_path=census,
        run_id=None,
    )


def _export(tmp_path: Path, scenarios: list[ExportScenario], fmt: str) -> Path:
    service = ComparisonExportService(storage=None, db_resolver=object())  # type: ignore[arg-type]
    out = tmp_path / "out"
    out.mkdir(exist_ok=True)
    return service.build(scenarios, "W", fmt, out)  # type: ignore[arg-type]


def _sheet_rows(path: Path, sheet: str) -> list[tuple]:
    return list(load_workbook(path, read_only=True)[sheet].iter_rows(values_only=True))


def test_xlsx_stacks_scenarios_on_one_tab_with_shared_census(tmp_path):
    scenarios = [_scenario(tmp_path, "Baseline", 3), _scenario(tmp_path, "Rich", 2)]

    path = _export(tmp_path, scenarios, "xlsx")

    assert load_workbook(path, read_only=True).sheetnames == [
        "Workforce_Snapshot",
        "Census",
        "Metadata",
    ]
    snapshot = _sheet_rows(path, "Workforce_Snapshot")
    assert snapshot[0][:3] == ("scenario_name", "scenario_id", "employee_id")
    assert [row[0] for row in snapshot[1:]] == ["Baseline"] * 6 + ["Rich"] * 4
    # The Studio ID replaces dbt's internal scenario_id ('default').
    assert snapshot[0].count("scenario_id") == 1
    assert snapshot[1][1] == "id-Baseline"
    census = _sheet_rows(path, "Census")
    assert census[0] == ("employee_id", "employee_hire_date")
    assert len(census) == 1 + 3


def test_xlsx_labels_census_when_scenarios_read_different_files(tmp_path):
    scenarios = [
        _scenario(tmp_path, "A", 2, census="a.parquet"),
        _scenario(tmp_path, "B", 1, census="b.parquet"),
    ]

    census = _sheet_rows(_export(tmp_path, scenarios, "xlsx"), "Census")

    assert census[0][:2] == ("scenario_name", "scenario_id")
    assert [row[0] for row in census[1:]] == ["A", "A", "B"]


def test_xlsx_uses_one_tab_per_scenario_over_row_limit(tmp_path, monkeypatch):
    monkeypatch.setattr(export_module, "EXCEL_MAX_DATA_ROWS", 8)
    scenarios = [
        _scenario(tmp_path, "Plan: A/B", 3),
        _scenario(tmp_path, "Plan: A?B", 3),
    ]

    path = _export(tmp_path, scenarios, "xlsx")

    workbook = load_workbook(path, read_only=True)
    assert workbook.sheetnames == ["Plan_ A_B", "Plan_ A_B (2)", "Census", "Metadata"]
    second = _sheet_rows(path, "Plan_ A_B (2)")
    assert second[0][:2] == ("scenario_name", "scenario_id")
    assert {row[0] for row in second[1:]} == {"Plan: A?B"}
    assert len(second) == 1 + 6


def test_xlsx_splits_a_single_oversized_scenario_by_year(tmp_path, monkeypatch):
    monkeypatch.setattr(export_module, "EXCEL_MAX_DATA_ROWS", 4)
    scenarios = [_scenario(tmp_path, "Small", 1), _scenario(tmp_path, "Big", 3)]

    path = _export(tmp_path, scenarios, "xlsx")

    assert load_workbook(path, read_only=True).sheetnames == [
        "Small",
        "Big_2025",
        "Big_2026",
        "Census",
        "Metadata",
    ]
    assert {row[2] for row in _sheet_rows(path, "Big_2026")[1:]} == {"e0", "e1", "e2"}


def test_xlsx_stacks_scenarios_whose_snapshot_columns_differ(tmp_path):
    older = _scenario(tmp_path, "Older", 1)
    conn = duckdb.connect(str(older.database_path))
    conn.execute("ALTER TABLE fct_workforce_snapshot DROP COLUMN current_compensation")
    conn.close()

    rows = _sheet_rows(
        _export(tmp_path, [_scenario(tmp_path, "Newer", 1), older], "xlsx"),
        "Workforce_Snapshot",
    )

    compensation = rows[0].index("current_compensation")
    assert [row[compensation] for row in rows[1:]] == [50000, 50000, None, None]


def test_unique_sheet_names_sanitize_truncate_and_dedupe():
    names = unique_sheet_names(["'Quoted'", "x" * 40, "X" * 40, "a[b]*c", ""])

    assert names[0] == "Quoted"
    assert names[1] == "x" * 31
    assert names[2] == "X" * 27 + " (2)"
    assert names[3] == "a_b__c"
    assert names[4] == "Scenario"
    assert all(len(name) <= 31 for name in names)


def test_hyper_has_all_scenarios_census_and_native_types(tmp_path):
    hyperapi = pytest.importorskip("tableauhyperapi")
    scenarios = [_scenario(tmp_path, "Baseline", 3), _scenario(tmp_path, "Rich", 2)]

    path = _export(tmp_path, scenarios, "hyper")

    with hyperapi.HyperProcess(
        hyperapi.Telemetry.DO_NOT_SEND_USAGE_DATA_TO_TABLEAU,
        parameters={"log_dir": str(tmp_path)},
    ) as hyper, hyperapi.Connection(hyper.endpoint, str(path)) as conn:
        counts = conn.execute_list_query(
            "SELECT scenario_name, COUNT(*) FROM workforce_snapshot "
            "GROUP BY scenario_name ORDER BY scenario_name"
        )
        census = conn.execute_scalar_query("SELECT COUNT(*) FROM census")
        columns = {
            column.name.unescaped: str(column.type)
            for column in conn.catalog.get_table_definition(
                hyperapi.TableName("workforce_snapshot")
            ).columns
        }
    assert counts == [["Baseline", 6], ["Rich", 4]]
    assert census == 3
    assert columns["simulation_year"] == "INT"
    assert columns["current_compensation"] == "NUMERIC(12, 2)"
    assert columns["snapshot_created_at"] == "TIMESTAMP_TZ"


# ---------------------------------------------------------------------------
# Endpoint
# ---------------------------------------------------------------------------


class _StubResolver:
    def __init__(self, paths: dict[str, Path]):
        self._paths = paths

    def resolve(self, workspace_id: str, scenario_id: str) -> ResolvedDatabasePath:
        path = self._paths.get(scenario_id)
        return ResolvedDatabasePath(path=path, source="scenario" if path else None)


@pytest.fixture
def env(tmp_path, monkeypatch):
    from fastapi.testclient import TestClient

    settings = api_config.APISettings(workspaces_root=tmp_path / "workspaces")
    monkeypatch.setattr(api_config, "settings", settings)
    storage = WorkspaceStorage(settings.workspaces_root)
    workspace = storage.create_workspace(WorkspaceCreate(name="My WS"), {})
    scenarios = []
    paths: dict[str, Path] = {}
    for name in ("Baseline", "Generous Match"):
        scenario = storage.create_scenario(workspace.id, ScenarioCreate(name=name))
        storage.update_scenario_status(workspace.id, scenario.id, "completed")
        paths[scenario.id] = _build_db(tmp_path / f"{scenario.id}.duckdb", 2)
        scenarios.append(scenario)

    app = create_app()
    app.dependency_overrides[
        analytics_router.get_comparison_export_service
    ] = lambda: ComparisonExportService(storage, _StubResolver(paths))
    return TestClient(app), storage, workspace, scenarios


def test_endpoint_downloads_xlsx(env):
    client, _, workspace, scenarios = env

    response = client.get(
        ENDPOINT.format(workspace_id=workspace.id),
        params={"scenarios": ",".join(s.id for s in scenarios)},
    )

    assert response.status_code == 200
    assert "My_WS_scenario_comparison_" in response.headers["content-disposition"]
    workbook = load_workbook(io.BytesIO(response.content), read_only=True)
    rows = list(workbook["Workforce_Snapshot"].iter_rows(values_only=True))
    assert {row[0] for row in rows[1:]} == {"Baseline", "Generous Match"}


def test_endpoint_allows_a_single_scenario(env):
    client, _, workspace, scenarios = env

    response = client.get(
        ENDPOINT.format(workspace_id=workspace.id),
        params={"scenarios": scenarios[0].id},
    )

    assert response.status_code == 200


def test_endpoint_rejects_unknown_and_unfinished_scenarios(env):
    client, storage, workspace, scenarios = env
    pending = storage.create_scenario(workspace.id, ScenarioCreate(name="Draft"))

    unknown = client.get(
        ENDPOINT.format(workspace_id=workspace.id), params={"scenarios": "nope"}
    )
    unfinished = client.get(
        ENDPOINT.format(workspace_id=workspace.id),
        params={"scenarios": f"{scenarios[0].id},{pending.id}"},
    )

    assert unknown.status_code == 404
    assert unfinished.status_code == 400


def test_endpoint_rejects_too_many_scenarios_and_bad_format(env):
    client, _, workspace, scenarios = env

    too_many = client.get(
        ENDPOINT.format(workspace_id=workspace.id),
        params={"scenarios": ",".join(f"s{i}" for i in range(7))},
    )
    bad_format = client.get(
        ENDPOINT.format(workspace_id=workspace.id),
        params={"scenarios": scenarios[0].id, "format": "csv"},
    )

    assert too_many.status_code == 400
    assert bad_format.status_code == 422
