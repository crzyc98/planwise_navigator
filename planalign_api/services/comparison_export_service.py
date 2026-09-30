"""Multi-scenario workforce export for the Cost Comparison page.

Produces an Excel workbook, Tableau ``.hyper`` extract, or ZIP of Parquet files
containing the full ``fct_workforce_snapshot`` of every selected scenario (labelled with
``scenario_name``/``scenario_id``) plus the cleaned census (``stg_census_data``)
the simulations read.

All scenario databases are ATTACHed read-only to one in-memory DuckDB
connection, so all writers stream from the same SQL.
"""

from __future__ import annotations

import re
import tempfile
import zipfile
from contextlib import contextmanager
from dataclasses import dataclass
from datetime import datetime, timezone
from pathlib import Path
from typing import Any, Iterator, Literal

import duckdb

from ..storage.workspace_storage import WorkspaceStorage
from .database_path_resolver import (
    DatabasePathResolver,
    create_api_database_path_resolver,
)

ExportFormat = Literal["xlsx", "hyper", "parquet"]

# Excel's hard per-sheet limit is 1,048,576 rows; one is the header.
EXCEL_MAX_DATA_ROWS = 1_048_575
_SHEET_NAME_MAX = 31
_INVALID_SHEET_CHARS = re.compile(r"[\[\]:*?/\\]")
_FETCH_BATCH_ROWS = 10_000
_CENSUS_PATH_KEYS = ("setup", "census_parquet_path")


class ComparisonExportError(ValueError):
    """A selected scenario cannot be exported."""


class HyperUnavailableError(RuntimeError):
    """The Tableau Hyper API is not installed on this platform."""


@dataclass(frozen=True)
class ExportScenario:
    scenario_id: str
    name: str
    database_path: Path
    census_path: str | None
    run_id: str | None


@dataclass(frozen=True)
class ExportTable:
    """One output sheet (Excel) or table (Hyper) and the query that fills it."""

    name: str
    query: str


class ComparisonExportService:
    def __init__(
        self,
        storage: WorkspaceStorage,
        db_resolver: DatabasePathResolver | None = None,
    ):
        self.storage = storage
        self.db_resolver = db_resolver or create_api_database_path_resolver(storage)

    def resolve_scenarios(
        self, workspace_id: str, scenario_ids: list[str], names: dict[str, str]
    ) -> list[ExportScenario]:
        """Pin each scenario to its selected run database and census input."""
        workspace = self.storage.get_workspace(workspace_id)
        scenarios = []
        for scenario_id in scenario_ids:
            resolved = self.db_resolver.resolve(workspace_id, scenario_id)
            if resolved.path is None:
                raise ComparisonExportError(
                    f"No results found for {names[scenario_id]}"
                )
            config = (
                self.storage.get_merged_config_for(workspace, scenario_id)
                if workspace
                else None
            )
            scenarios.append(
                ExportScenario(
                    scenario_id=scenario_id,
                    name=names[scenario_id],
                    database_path=resolved.path,
                    census_path=_nested_str(config, _CENSUS_PATH_KEYS),
                    run_id=resolved.run_id,
                )
            )
        return scenarios

    def build(
        self,
        scenarios: list[ExportScenario],
        workspace_name: str,
        export_format: ExportFormat,
        output_dir: Path,
    ) -> Path:
        """Write the export into ``output_dir`` and return the file path."""
        extension = "zip" if export_format == "parquet" else export_format
        path = output_dir / f"scenario_comparison.{extension}"
        with _attached(scenarios) as conn:
            _create_metadata_table(conn, scenarios, workspace_name)
            if export_format == "hyper":
                _write_hyper(conn, _hyper_tables(scenarios), path)
            elif export_format == "parquet":
                _write_parquet_archive(conn, _hyper_tables(scenarios), path)
            else:
                _write_xlsx(conn, _excel_sheets(conn, scenarios), path)
        return path


# ---------------------------------------------------------------------------
# Query construction
# ---------------------------------------------------------------------------


_SNAPSHOT = "fct_workforce_snapshot"
_CENSUS = "stg_census_data"
# Current snapshots carry dbt's internal scenario_id ('default'); the Studio
# scenario ID replaces it rather than colliding with it.
_LABEL_COLUMNS = ("scenario_name", "scenario_id")


@contextmanager
def _attached(scenarios: list[ExportScenario]) -> Iterator[duckdb.DuckDBPyConnection]:
    """ATTACH each run database and expose ``s<i>_<table>`` labelled views."""
    conn = duckdb.connect()
    try:
        # TIMESTAMPTZ columns render in UTC in all writers.
        conn.execute("SET TimeZone = 'UTC'")
        for index, scenario in enumerate(scenarios):
            conn.execute(
                f"ATTACH {_literal(str(scenario.database_path))} "
                f"AS {_alias(index)} (READ_ONLY)"
            )
            for table in (_SNAPSHOT, _CENSUS):
                _create_labelled_view(conn, index, scenario, table)
        yield conn
    finally:
        conn.close()


def _create_labelled_view(
    conn: duckdb.DuckDBPyConnection, index: int, scenario: ExportScenario, table: str
) -> None:
    columns = conn.execute(
        "SELECT column_name FROM duckdb_columns() "
        "WHERE database_name = ? AND schema_name = 'main' AND table_name = ? "
        "ORDER BY column_index",
        [_alias(index), table],
    ).fetchall()
    if not columns:
        raise ComparisonExportError(f"{scenario.name} has no {table} table to export")
    select_list = ", ".join(
        '"' + name.replace('"', '""') + '"'
        for (name,) in columns
        if name not in _LABEL_COLUMNS
    )
    conn.execute(
        f"CREATE TEMP VIEW {_view(index, table)} AS "
        f"SELECT {_literal(scenario.name)} AS scenario_name, "
        f"{_literal(scenario.scenario_id)} AS scenario_id, "
        f"{index} AS _scenario_order, {select_list} "
        f"FROM {_alias(index)}.main.{table}"
    )


def _alias(index: int) -> str:
    return f"s{index}"


def _view(index: int, table: str) -> str:
    return f"{_alias(index)}_{table}"


def _literal(value: str) -> str:
    return "'" + value.replace("'", "''") + "'"


def _labelled_union(
    indices: list[int], table: str, order_by: str, where: str = ""
) -> str:
    """Stack ``table`` across scenarios; BY NAME tolerates schema drift between runs."""
    parts = " UNION ALL BY NAME ".join(
        f"SELECT * FROM {_view(index, table)}" for index in indices
    )
    return (
        f"SELECT * EXCLUDE (_scenario_order) FROM ({parts}) {where} "
        f"ORDER BY _scenario_order, {order_by}"
    )


def _snapshot_query(indices: list[int], where: str = "") -> str:
    return _labelled_union(indices, _SNAPSHOT, "simulation_year, employee_id", where)


def _census_query(scenarios: list[ExportScenario]) -> str:
    """One unlabelled census when every scenario read the same input file."""
    census_paths = {scenario.census_path for scenario in scenarios}
    shared_census = len(census_paths) == 1 and None not in census_paths
    if shared_census:
        return (
            f"SELECT * EXCLUDE (scenario_name, scenario_id, _scenario_order) "
            f"FROM {_view(0, _CENSUS)} ORDER BY employee_id"
        )
    return _labelled_union(list(range(len(scenarios))), _CENSUS, "employee_id")


def _create_metadata_table(
    conn: duckdb.DuckDBPyConnection,
    scenarios: list[ExportScenario],
    workspace_name: str,
) -> None:
    rows = [
        ("exported_at_utc", datetime.now(timezone.utc).isoformat(timespec="seconds")),
        ("workspace", workspace_name),
    ]
    for scenario in scenarios:
        rows.extend(
            [
                ("scenario_name", scenario.name),
                ("scenario_id", scenario.scenario_id),
                ("run_id", scenario.run_id or ""),
                ("census_path", scenario.census_path or ""),
            ]
        )
    conn.execute("CREATE TEMP TABLE export_metadata (field VARCHAR, value VARCHAR)")
    conn.executemany("INSERT INTO export_metadata VALUES (?, ?)", rows)


_METADATA_QUERY = "SELECT field, value FROM export_metadata ORDER BY rowid"


# ---------------------------------------------------------------------------
# Excel layout
# ---------------------------------------------------------------------------


def _excel_sheets(
    conn: duckdb.DuckDBPyConnection, scenarios: list[ExportScenario]
) -> list[ExportTable]:
    """One combined snapshot tab when it fits; otherwise one tab per scenario."""
    counts = [_row_count(conn, index) for index in range(len(scenarios))]
    if sum(counts) <= EXCEL_MAX_DATA_ROWS:
        snapshot = [
            ExportTable(
                "Workforce_Snapshot", _snapshot_query(list(range(len(scenarios))))
            )
        ]
    else:
        snapshot = [
            sheet
            for index, scenario in enumerate(scenarios)
            for sheet in _scenario_sheets(conn, index, scenario, counts[index])
        ]
    sheets = snapshot + [
        ExportTable("Census", _census_query(scenarios)),
        ExportTable("Metadata", _METADATA_QUERY),
    ]
    names = unique_sheet_names([sheet.name for sheet in sheets])
    return [ExportTable(name, sheet.query) for name, sheet in zip(names, sheets)]


def _row_count(conn: duckdb.DuckDBPyConnection, index: int) -> int:
    row = conn.execute(
        f"SELECT COUNT(*) FROM {_alias(index)}.main.{_SNAPSHOT}"
    ).fetchone()
    return int(row[0]) if row else 0


def _scenario_sheets(
    conn: duckdb.DuckDBPyConnection, index: int, scenario: ExportScenario, rows: int
) -> list[ExportTable]:
    """A scenario tab, split by simulation year if it alone exceeds the limit."""
    single = [index]
    if rows <= EXCEL_MAX_DATA_ROWS:
        return [ExportTable(scenario.name, _snapshot_query(single))]
    years = conn.execute(
        f"SELECT DISTINCT simulation_year FROM {_alias(index)}.main.{_SNAPSHOT} "
        "ORDER BY simulation_year"
    ).fetchall()
    base = sanitize_sheet_name(scenario.name)[: _SHEET_NAME_MAX - 5]
    return [
        ExportTable(
            f"{base}_{year}",
            _snapshot_query(single, f"WHERE simulation_year = {int(year)}"),
        )
        for (year,) in years
    ]


def sanitize_sheet_name(name: str) -> str:
    """Strip characters Excel rejects and cap at 31 characters."""
    cleaned = _INVALID_SHEET_CHARS.sub("_", name).strip().strip("'").strip()
    return cleaned[:_SHEET_NAME_MAX] or "Scenario"


def unique_sheet_names(names: list[str]) -> list[str]:
    """Sanitize names and suffix ``(2)``, ``(3)`` … on case-insensitive clashes."""
    seen: set[str] = set()
    result = []
    for raw in names:
        base = sanitize_sheet_name(raw)
        candidate, counter = base, 2
        while candidate.lower() in seen:
            suffix = f" ({counter})"
            candidate = base[: _SHEET_NAME_MAX - len(suffix)] + suffix
            counter += 1
        seen.add(candidate.lower())
        result.append(candidate)
    return result


def _write_xlsx(
    conn: duckdb.DuckDBPyConnection, sheets: list[ExportTable], path: Path
) -> None:
    from openpyxl import Workbook

    # write_only streams rows to disk instead of holding every cell in memory.
    workbook = Workbook(write_only=True)
    for sheet in sheets:
        _append_sheet(workbook, conn, sheet)
    workbook.save(path)


def _append_sheet(
    workbook: Any, conn: duckdb.DuckDBPyConnection, sheet: ExportTable
) -> None:
    from openpyxl.cell import WriteOnlyCell
    from openpyxl.styles import Font
    from openpyxl.utils import get_column_letter

    worksheet = workbook.create_sheet(sheet.name)
    worksheet.freeze_panes = "A2"
    cursor = conn.execute(sheet.query)
    headers = [column[0] for column in cursor.description or []]
    for position, header in enumerate(headers, start=1):
        worksheet.column_dimensions[get_column_letter(position)].width = max(
            len(header) + 2, 12
        )
    header_cells = []
    for header in headers:
        cell = WriteOnlyCell(worksheet, value=header)
        cell.font = Font(bold=True)
        header_cells.append(cell)
    worksheet.append(header_cells)
    while batch := cursor.fetchmany(_FETCH_BATCH_ROWS):
        for row in batch:
            worksheet.append([_excel_value(value) for value in row])


def _excel_value(value: Any) -> Any:
    """openpyxl rejects tz-aware datetimes; values are already rendered in UTC."""
    if isinstance(value, datetime) and value.tzinfo is not None:
        return value.replace(tzinfo=None)
    return value


# ---------------------------------------------------------------------------
# Tableau Hyper
# ---------------------------------------------------------------------------


def _hyper_tables(scenarios: list[ExportScenario]) -> list[ExportTable]:
    return [
        ExportTable("workforce_snapshot", _snapshot_query(list(range(len(scenarios))))),
        ExportTable("census", _census_query(scenarios)),
        ExportTable("metadata", _METADATA_QUERY),
    ]


def _write_hyper(
    conn: duckdb.DuckDBPyConnection, tables: list[ExportTable], path: Path
) -> None:
    """Stage each table as Parquet, then let Hyper ingest it natively (types kept)."""
    try:
        from tableauhyperapi import (
            Connection,
            CreateMode,
            HyperProcess,
            Telemetry,
            escape_name,
            escape_string_literal,
        )
    except ImportError as exc:
        raise HyperUnavailableError(
            "Tableau Hyper export is not available on this platform"
        ) from exc

    with tempfile.TemporaryDirectory(dir=path.parent) as staging:
        staging_dir = Path(staging)
        with HyperProcess(
            Telemetry.DO_NOT_SEND_USAGE_DATA_TO_TABLEAU,
            parameters={"log_dir": str(staging_dir)},
        ) as hyper, Connection(
            hyper.endpoint, str(path), CreateMode.CREATE_AND_REPLACE
        ) as hyper_conn:
            for table in tables:
                parquet = staging_dir / f"{table.name}.parquet"
                conn.execute(
                    f"COPY ({table.query}) TO {_literal(str(parquet))} (FORMAT PARQUET)"
                )
                hyper_conn.execute_command(
                    f"CREATE TABLE {escape_name(table.name)} AS "
                    f"(SELECT * FROM external({escape_string_literal(str(parquet))}))"
                )


def _write_parquet_archive(
    conn: duckdb.DuckDBPyConnection, tables: list[ExportTable], path: Path
) -> None:
    """Export native types to one Parquet file per dataset in a ZIP download."""
    with tempfile.TemporaryDirectory(dir=path.parent) as staging:
        with zipfile.ZipFile(path, "w", compression=zipfile.ZIP_STORED) as archive:
            for table in tables:
                parquet = Path(staging) / f"{table.name}.parquet"
                conn.execute(
                    f"COPY ({table.query}) TO {_literal(str(parquet))} (FORMAT PARQUET)"
                )
                archive.write(parquet, arcname=parquet.name)


def _nested_str(config: dict[str, Any] | None, keys: tuple[str, ...]) -> str | None:
    value: Any = config
    for key in keys:
        if not isinstance(value, dict):
            return None
        value = value.get(key)
    return str(value) if value else None
